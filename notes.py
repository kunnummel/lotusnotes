#! python3
#Author: Jis Joseph Kunnummel - www.kunnummel.com
class UIBack:
  s=None
  db=None
  mdb=None
  mvw=None
  mnav=None
  vw=None
  vn=None
  ve=None
  doc=None
  dc=None
  dcs=None
  user=None
  _evaldoc=None
  def init():
    if(UIBack.user==None):
      try:
        from win32com.client import Dispatch
        UIBack.s=Dispatch('Lotus.NotesSession')
        UIBack.user=UIBack.s.username
      except:
        try:
          UIBack.s.initialize()
          UIBack.user=UIBack.s.username
        except Exception as e:
          print(e)
      finally:
        return UIBack.user
    else:
      return UIBack.user
  def evalformula(doc=None,formula='@username'):
    if(UIBack.init()==None):
      return
    try:
      if(UIBack._evaldoc==None):
        UIBack._evaldoc=UIBack.s.urldatabase.createdocument()
      return UIBack.s.evaluate(formula,doc if doc else UIBack._evaldoc)
    except Exception as e:
      print(e)

  def resolve(arg):
    if(UIBack.init()==None):
      return (None,None)
    try:
      urlnote=UIBack.s.resolve(arg)
      import re
      tp=urlnote.notesurl.split('?Open',1)[-1]
    except:
      return (None,None)    
    return (urlnote,urlnote.parentdatabase if( tp== 'Document') else urlnote if tp=='Database' else urlnote.parent,tp)
    
  def grab():
    if(UIBack.init()==None):
      return
    if(input("Retrieve backend data (COM) for the currently active 'Notes UI'? - (y)/n => ")=='n'):
      return
    UI.grab(False)
    UIBack.db=UI.db
    if(UIBack.db==None):
      return
    UIBack.db=UIBack.s.getdatabase(UIBack.db.server,UIBack.db.filepath) if UIBack.db else None
    if(UI.vw):
      UIBack.vw=UIBack.db.getView(UI.vw.name)
      UIBack.vw.autoupdate=False
      UIBack.vn=UIBack.vw.createviewnav()
      if(UI.dc.count>0 and input(f'Do you also would like to get the {UI.dc.count} documents selected at the UI view? (n)/y >') =='y'):
        UIBack.dc=UIBack.vw.getalldocumentsbykey('\n\n',True)
        UIBack.dcs={}
        Loop.runondocs(UI.dc,lambda x:UIBack.dc.adddocument(UIBack.db.getdocumentbyid(x.noteid)))
        def aa(x):
          UIBack.vn.gotoentry(x)
          return ({x.noteid:list(UIBack.vn.getcurrent().columnvalues)})
        UIBack.dcs=Loop.iterrecords(UIBack.dc,aa)
    try:
      UIBack.doc=UIBack.db.getdocumentbyid(UIBack.vw.caretnoteid)
    except:
      UIBack.doc=UIBack.db.getdocumentbyid(UI.doc.noteid) if UI.doc else None
    if(UIBack.doc and UIBack.vn):
      UIBack.vn.gotoentry(UIBack.doc)
      UIBack.ve=UIBack.vn.getcurrent()
    print(f'''
UIBack.db > {(UIBack.db.server,UIBack.db.filepath,UIBack.db.title,UIBack.db.replicaid) if UIBack.db else 'n/a'}
UIBack.vw > {(UIBack.vw.name,UIBack.vw.aliases,UIBack.vw.entrycount,UIBack.vw.universalid,{k:v for a in [{x.position:x.title+(' (H)' if x.ishidden else '')} for x in UIBack.vw.columns if x.title!='' or x.ishidden] for k,v in a.items()}) if UIBack.vw else 'n/a'}
UIBack.vn/.ve> Position: {(UIBack.ve.getposition('.'),UI.ws.currentview.caretcategory) if UIBack.ve else 'n/a'} 
UIBack.doc> {(UIBack.doc.getitemvalue("form")[0],UIBack.doc.universalid,UIBack.doc.noteid,UIBack.doc.size,UIBack.doc.hasembedded,{a[0]+1:str(a[1]) for a in enumerate(UIBack.ve.columnvalues)} if UIBack.ve else ('new-doc' if UIBack.doc.isnewnote else 'items-'+str(len(UIBack.doc.items)))) if UIBack.doc else 'n/a'}
UIBack.dc/.dcs > {UIBack.dc.count if UIBack.dc else 'n/a'}
''')
  def verifydoc(doc):
    return doc.parentdatabase.parent==UIBack.s
class UI:
  ws=None
  s=None
  db=None
  mdb=None
  mvw=None
  mnav=None
  vw=None
  vn=None
  ve=None
  doc=None
  dc=None
  dcs=None
  user=None
  def init():
    try:
      if(UI.user==None):
        from win32com.client import Dispatch
        UI.ws=Dispatch('Notes.NotesUIWorkspace')
        UI.s=Dispatch('Notes.NotesSession')
        UI.user=UI.s.username
    except: pass
    finally: return UI.user 
  def grab(warn=True):
    if(UI.init()==None):
      return
    if(warn):
      if(input("Retrieve currently active 'Notes UI' data? - (y)/n => ")=='n'):
        return
    UI.user=UI.s.username
    while True:
      UI.db=UI.ws.currentdatabase
      if(UI.db or input("No active Database objects found. Try again (y)/n => ")=='n'):
          break
    if(UI.db==None):
      return
    UI.vw=UI.ws.currentview
    UI.dc=UI.vw.documents if UI.vw else None
    UI.db=UI.db.database if UI.db else None
    try:
      UI.doc=UI.db.getdocumentbyid(UI.vw.caretnoteid)
    except:
      UI.doc=UI.ws.currentdocument
      UI.doc=UI.doc.document if UI.doc else None
    if(UI.vw):
      UI.vw=UI.vw.view
      UI.vw.autoupdate=False
    UI.vn=UI.vw.createviewnav if UI.vw else None
    UI.ve=(UI.vn.getentry(UI.doc) if UI.doc else UI.vn.getprev(UI.vw.getentrybykey(UI.ws.currentview.caretcategory,True))) if UI.vn else None
    if(UI.dc):
      UI.dcs=Loop.iterrecords(UI.dc,lambda x:({x.noteid:list(UI.vn.getentry(x).columnvalues)}))
    print(f'''
UI.db > {(UI.db.server,UI.db.filepath,UI.db.title,UI.db.replicaid) if UI.db else 'n/a'}
UI.vw > {(UI.vw.name,UI.vw.aliases,UI.vw.entrycount,UI.vw.universalid,{k:v for a in [{x.position:x.title+(' (H)' if x.ishidden else '')} for x in UI.vw.columns if x.title!='' or x.ishidden] for k,v in a.items()}) if UI.vw else 'n/a'}
UI.vn/.ve> Position: {UI.ve.getposition('.') if UI.ve else 'n/a'} Cursor Category - {UI.ws.currentview.caretcategory} 
UI.doc> {(UI.doc.form[0],UI.doc.universalid,UI.doc.noteid,UI.doc.size,UI.doc.hasembedded,{a[0]+1:str(a[1]) for a in enumerate(UI.ve.columnvalues)} if UI.ve else ('new-doc' if UI.doc.isnewnote else 'items-'+str(len(UI.doc.items)))) if UI.doc else 'n/a'}
UI.dc/.dcs > {UI.dc.count if UI.dc else 'n/a'}
''')
  def resolve(arg):
    if(UI.init()==None):
      return (None,None)
    try:
      urlnote=UI.s.resolve(arg)
      tp=urlnote.notesurl.split('?Open',1)[-1]
    except:
      return (None,None)
    return (urlnote,urlnote.parentdatabase if( tp== 'Document') else urlnote if tp=='Database' else urlnote.parent,tp)

  def locatedoc(doc=None):
    doc=doc if doc else UI.doc
    if(doc):
      if(not UI.verifydoc(doc)):
        print('Backend document is not allowed')
        return
      vw=UI.ws.currentview
      vw.selectdocument(doc)
      print(vw.caretcategory if vw.caretnoteid==doc.noteid else 'Document could not be located')
  def showwindows(title='HCL Notes',nameit=False):
    try:
      import win32gui,win32process
      listnames={}
      def showMe(hs,*args):
        st=win32gui.GetWindowText(hs)
        listnames.setdefault(st, []).append([hs,win32process.GetWindowThreadProcessId(hs)])
        if(not None==title and title in st):        
          win32gui.ShowWindow(hs,1) if input(f'Show Window - {st} - (Process - {hs}) y/(n) >')=='y' else None
      win32gui.EnumWindows(showMe,None)
      if(nameit):
        return dict(sorted(listnames.items()))
    except:print('error')
  def checkmail(cnt=5):
    UI.init()
    if(UI.s and (not UI.mdb)):
      UI.mdb=UI.s.getdatabase('','')
      if(not UI.mdb.isopen):
        UI.mdb.openmail
    if (not UI.mdb):
      return
    if(not UI.mnav):
      UI.mvw=UI.mdb.getview('($Inbox)')
      if(UI.mvw==None):
        return
      UI.mnav=UI.mvw.createviewnav
    ve=UI.mnav.getlast
    for x in range(cnt):
      if(ve):
        doc=ve.document
        print(doc.posteddate[0],' - ', doc.getitemvalue('from')[0],'\n',doc.subject[0],'\n',doc.getitemvalue('$abstract')[0],'\n','-'*100,'\n')
        ve=UI.mnav.getprev(ve)
  def select(vw=None,category=None):
    if (vw==None and UI.vw):
      vw=UI.vw
    if(vw==None):
      return (0,None)
    print('Switch to the Notes Window to select the Document(s)')
    dc=UI.ws.PickListCollection(3,True,vw.parent.server,vw.parent.filepath,vw.name,'Select Document(s)','Pick one or more documents from the view',category)
    return(dc.getfirstdocument,dc if dc.count>1 else None)
  def browsefiles(filter=None,save=False):
    return (UI.ws.SaveFileDialog(False,'Save As',filter,'C:\\')) if save else (UI.ws.OpenFileDialog(True,'Open',filter,'C:\\'))
  def verifydoc(doc):
    return doc.parentdatabase.parent==UI.s
class Utils:
  def getattachments(doc,details=False):
    att={} if details else []
    if(doc.hasembedded):
      for x in doc.items:
        if(x.type==1084):
          for y in x.values:
            f=doc.getattachment(y)
            att.update({f.name:(f,f.filesize,f.source)}) if details else att.append(f)
    return {doc.noteid:att} if details else att

  def dataframe(dcs=None):
    import pandas as pd
    df = pd.DataFrame(dcs if dcs else list(UI.dcs))
    return df.transpose()
  def dictupdatecounter(dictobj,key='',start=1,incr=1):
    dictobj.update({key:dictobj.setdefault(key,start)+incr})
  def flatten(lst,result=[]):
    for x in lst:
      if(isinstance(x,(list,tuple,set))):
        Utils.flatten(x,result)
      elif(isinstance(x,dict)):
        Utils.flatten((tuple(x.keys()),tuple(x.values())),result)
      else:
        result.append(str(x))
    return result
class Loop:
  def dbdirectory(server='',type=1247):
    dbr=UI.s.getdbdirectory(server)
    db=dbr.getfirstdatabase(type)
    dbrdict={}
    while db:
      dbrdict.update({db.filepath:(db.title,db)})
      db=dbr.getnextdatabase
    return dbrdict
  def export(dc,f,*args,file='C:/temp/tempfile.xlsx',**kwargs):
    from openpyxl import Workbook
    wb=Workbook()
    wbs=wb.active
    def aa(x):
      for xx in x:
        wbs.append(['\n'.join([str(z) for z in y]) if isinstance(y,tuple) else str(y) for y in xx])
      wb.save(file)
    Loop.runondocs(dc,f,*args,callback=aa,**kwargs)
    wb.close()
    print ("-- Generated file ",file) 
  def exportasjson(dc,f,*args,**kwargs):
    wbs=[]
    def aa(x):
      for xx in x:
        wbs.append(['\n'.join([str(z) for z in y]) if isinstance(y,tuple) else str(y) for y in xx])
    
    Loop.runondocs(dc,f,*args,callback=aa,**kwargs)
    return wbs

  def exportdocfields(dc,fields,file='C:/temp/tempfile.xlsx'):
    from openpyxl import Workbook
    wb=Workbook()
    wbs=wb.active
    fields=fields.split(',') if isinstance(fields,str) else fields
    wbs.append(fields)
    def aa(x,w=wbs,f=fields):
      try:
        w.append([''.join(x.getitemvalue(y)) for y in f])
      except Exception as e:
        print('ERR> ',e)
    Loop.runondocs(dc,aa)
    wb.save(file)
    wb.close()
    print ("-- Generated file ",file)      
  def docvaluescount(dc=None,itms=['form']):
    aa={}
    if(dc==None):
      dc=UI.vw
    Loop.runondocs(dc,lambda x:[aa.update({x.getitemvalue(y)[0]:aa.dictobj.setdefault(x.getitemvalue(y)[0],1)+1}) for y in itms])
    if(len(aa)>0):
      return(aa)
    return None
  def docitemsdict(itms,values=False,nameonly=True):
    aa={'text':[],'number':[],'rt':[],'date':[],'name':[],'att':[],'others':[]}
    for x in itms:
      if(x.type in [1280,1282,21]):
        iname='text'
      elif(x.type==768):
        iname='number'
      elif(x.type==1024):
        iname='date'
      elif(x.type==1084):
        iname='att'
      elif(x.type in [1,4,6,7,8,25,1090]):
        iname='rt'
      elif(x.type in [1074,1075,1076]):
        iname='name'
      else:
        iname='others'
      if nameonly:
        bb=x.name
      else:
        bb= [x.name,x.valuelength,x.type]
        if(values):
          bb.append(x.text)
      aa.get(iname).append(bb)
    return(aa)

  def iterrecords(coll,func,*args,entries=False,startrec=None,startat=1,forward=True,details=False,result=None,**kwargs):
    if(coll):
      if(startrec):
        rec=startrec
        idc=1
      else:
        idc=startat
        try: 
          if(entries):
            rec=coll.getlastentry() if startat==-1 else coll.getfirstentry() if startat==1 else coll.getnthentry(startat) 
          else:
            rec=coll.getlastdocument() if startat==-1 else coll.getfirstdocument() if startat==1 else coll.getnthdocument(startat)
        except:
          if(entries):
            rec=coll.getlastentry if startat==-1 else coll.getfirstentry
          else:
            rec=coll.getlastdocument if startat==-1 else coll.getfirstdocument     
      while rec:
        if(entries):
          trec=coll.getnextentry(rec) if forward else coll.getpreventry(rec) 
        else:
          trec=coll.getnextdocument(rec) if forward else coll.getprevdocument(rec) 
        cb=func(rec,*args,**kwargs) if func!=None else rec
        aa=(idc,rec,cb) if details else cb
        result.append(aa) if result!=None else None
        yield aa
        idc=idc+1 if forward else idc-1
        if(Loop.stop(rec.noteid,idc)):
          break
        rec=trec
  def iternotes(nc,func,*args):
    n1=nc.getfirstnoteid
    while n1!='':
      n2=nc.getnextnoteid(n1)
      yield func(n1,*args)
      n1=n2
  def iternav(nv,func,*args,start=None,forward=True,samelevel=False):
    n1=nv.getfirst if start==None else nv.getentry(start)
    while n1:
      if(samelevel):
        n2=nv.getnextsibling(n1) if forward else nv.getprevsibling(n1)
      else:
        n2=nv.getnext(n1) if forward else nv.getprev(n1)
      yield func(n1,*args)
      n1=n2
  def runondocs(dc,func,*args,cbfunc=None,**kwargs):
    if(dc and func):
      try:
        idc=[]
        d1=dc.getfirstdocument()
      except:
        d1=dc.getfirstdocument
      while d1:
        try:
          d2=dc.getnextdocument(d1)
          if(Loop.stop(d1.noteid,idc)):
            break
          idc.append(func(d1,*args,**kwargs))
        except Exception as e:
          print('ERR> ',d1.noteid,idc,e)	
          if(input('Do you wish to break it? (y)/n >')!='n'):
            break
        d1=d2
      if(cbfunc):
        cbfunc(idc)
  def itercolumnvalues(vw=None,cols=[]):
    pass
  async def async_runondocs(dc,func,*args,cbfunc=None,**kwargs):
    if(dc and func):
      try:
        idc=[]
        d1=dc.getfirstdocument()
      except:
        d1=dc.getfirstdocument
      while d1:
        try:
          d2=dc.getnextdocument(d1)
          if(Loop.stop(d1.noteid,len(idc))):
            break
          idc.append(await func(d1,*args,**kwargs))
        except Exception as e:
          print('ERR> ',d1.noteid,idc,e)	
          if(input('Do you wish to break it? (y)/n >')!='n'):
            break
        d1=d2
      if(cbfunc):
        cbfunc(idc)
  def runonnotes(nc,func,*args,cbfunc=None,**kwargs):
    if(nc and func):
      idc=[]
      n1=nc.getfirstnoteid
      while n1:
        try:
          n2=nc.getnextnoteid(n1)
          if(Loop.stop(n1)):
            break
          idc.append(func(n1,*args,**kwargs))
        except Exception as e:
          print('ERR> ',n1,e)	
          if(input('Do you wish to break it? (y)/n >')!='n'):
            break
        n1=n2
      if(cbfunc):
        cbfunc(idc)
  def runonentries(ec,func,*args,cbfunc=None,**kwargs):
    if(ec and func):
      try:
        idc=[]
        e1=ec.getfirstentry()
      except:
        e1=ec.getfirstentry
      while e1:
        try:
          e2=ec.getnextentry(e1)
          if(Loop.stop(e1,idc)):
            break
          idc.append(func(e1,*args,**kwargs))
        except Exception as e:
          print('ERR> ',e1,idc,e)
        e1=e2
      if(cbfunc):
        cbfunc(idc)
  def stop(*args,key='ctrl+c'):
    try:
      import keyboard
      if(keyboard.is_pressed(key)):
        if(input(f'>> Do you wish to interrupt? {args} (n)/Y > ')=='Y'):
          print('-- Stopped at user request at ',args)
          return True
    except Exception as e:
      print('ERR> ',e)
      return True
  def getdocids(dc,unid=False,entries=False):
    selids=[]
    if(entries):
      Loop.runonentries(dc,lambda e1:selids.append(e1.document.universalid if unid else e1.document.noteid))
    else:
      Loop.runondocs(dc,lambda d1:selids.append(d1.universalid if unid else d1.noteid))
    return(selids)

  def getiddocs(selids=[],db=UI.db):
    seldocs={}
    if (db==None):
      db=UI.db
    for id in selids:
      seldocs.update({id:db.getdocumentbyunid(id) if len(id)==32 else db.getdocumentbyid(id)})
    return(seldocs)

  def comprops(obj=[],prop=['name']):
    if(not isinstance(obj,(list,tuple))):
      obj=[obj]
    if(not isinstance(prop,(list,tuple))):
      prop=[prop]
    _obj=[]
    for x in obj:
      for y in prop:
        try:
          _obj.append(x.__getattr__(y))
        except Exception as e:
          print('ERR> ',x,y,e)
      if(Loop.stop(x,y)):
        return _obj
    return _obj

def main():
  try:
    if(input('Access Notes Workspace? y/(n) > ')=='y'):
      UI.init()
    if(input('Access Notes COM Session? y/(n) > ')=='y'):
      UIBack.init()
  except Exception as e:
    print('ERR> ',e)
  finally:
    print(f'Global>> UI.ws>{'Workspace' if UI.ws else None}, UI.s>{UI.s.username if UI.s else None}, UIBack.s>{UIBack.user}, Loop, Utils')

if __name__ == "__main__":
  main()